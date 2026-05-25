import hashlib
import os
import secrets
from datetime import date, datetime
from functools import wraps
from io import BytesIO

import pandas as pd
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS

from database import get_connection, init_db
from forecast import generate_demo_data, get_forecast

app = Flask(__name__)
CORS(app)

_tokens: dict[str, dict] = {}

init_db()


# ─────────────────────────────────────────
# Serve frontend
# ─────────────────────────────────────────

@app.get("/")
def index():
    path = os.path.join(os.path.dirname(__file__), "..", "frontend", "index.html")
    return send_file(os.path.abspath(path))


# ─────────────────────────────────────────
# Auth helpers
# ─────────────────────────────────────────

def require_auth(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        token = request.headers.get("Authorization", "").removeprefix("Bearer ").strip()
        if not token or token not in _tokens:
            return jsonify({"error": "Ej autentiserad"}), 401
        return f(*args, **kwargs)
    return wrapper


def _days_param(default: int = 7) -> int:
    """Read ?days= query param, clamp to [1, 90]."""
    try:
        return max(1, min(90, int(request.args.get("days", default))))
    except (ValueError, TypeError):
        return default


# ─────────────────────────────────────────
# Auth endpoints
# ─────────────────────────────────────────

@app.post("/api/login")
def login():
    body     = request.get_json(silent=True) or {}
    email    = body.get("email", "").strip()
    password = body.get("password", "")
    if not email or not password:
        return jsonify({"error": "Email och lösenord krävs"}), 400
    pw_hash  = hashlib.sha256(password.encode()).hexdigest()
    conn     = get_connection()
    user     = conn.execute(
        "SELECT id FROM users WHERE email = ? AND password_hash = ?", (email, pw_hash)
    ).fetchone()
    conn.close()
    if not user:
        return jsonify({"error": "Felaktiga inloggningsuppgifter"}), 401
    token = secrets.token_urlsafe(32)
    _tokens[token] = {"email": email, "user_id": user["id"]}
    return jsonify({"token": token, "email": email})


@app.post("/api/logout")
@require_auth
def logout():
    token = request.headers.get("Authorization", "").removeprefix("Bearer ").strip()
    _tokens.pop(token, None)
    return jsonify({"success": True})


# ─────────────────────────────────────────
# Sales upload  (with auto stock deduction)
# ─────────────────────────────────────────

@app.post("/api/upload")
@require_auth
def upload():
    if "file" not in request.files:
        return jsonify({"error": "Ingen fil uppladdad"}), 400
    f    = request.files["file"]
    name = (f.filename or "").lower()
    try:
        if name.endswith(".csv"):
            df = pd.read_csv(f, encoding="utf-8-sig")
        elif name.endswith((".xlsx", ".xls")):
            df = pd.read_excel(f)
        else:
            return jsonify({"error": "Stödjer bara CSV och Excel"}), 400

        df.columns = df.columns.str.lower().str.strip()
        VARA_NAMES  = {"vara", "produkt", "product", "item", "artikel"}
        DATUM_NAMES = {"datum", "date", "dag", "tid"}
        ANTAL_NAMES = {"antal", "quantity", "qty", "amount", "försäljning", "sold"}

        col_vara  = next((c for c in df.columns if c in VARA_NAMES), None)
        col_datum = next((c for c in df.columns if c in DATUM_NAMES), None)
        col_antal = next((c for c in df.columns if c in ANTAL_NAMES), None)
        missing   = [n for n, c in [("vara", col_vara), ("datum", col_datum), ("antal", col_antal)] if not c]
        if missing:
            return jsonify({"error": f"Saknar kolumn(er): {', '.join(missing)}"}), 400

        # Group rows by product  →  {name: [(date, qty), ...]}
        upload_map: dict[str, list[tuple[str, int]]] = {}
        for _, row in df.iterrows():
            try:
                pname    = str(row[col_vara]).strip()
                date_val = str(row[col_datum]).strip()
                qty      = int(float(str(row[col_antal])))
                if not pname or pname == "nan":
                    continue
                upload_map.setdefault(pname, []).append((date_val, qty))
            except (ValueError, TypeError):
                continue

        conn          = get_connection()
        cur           = conn.cursor()
        inserted      = 0
        auto_deducted = []

        for pname, entries in upload_map.items():
            # Ensure product row exists
            cur.execute("INSERT OR IGNORE INTO products (name) VALUES (?)", (pname,))
            prod = cur.execute(
                "SELECT id, current_stock, stock_initialized, stock_sync_date FROM products WHERE name = ?",
                (pname,),
            ).fetchone()

            max_upload_date = max(d for d, _ in entries)

            # Insert all sales rows
            for date_val, qty in entries:
                cur.execute(
                    "INSERT INTO sales (product_id, date, quantity) VALUES (?, ?, ?)",
                    (prod["id"], date_val, qty),
                )
                inserted += 1

            # Auto-deduct from stock (only for dates after the sync point)
            if prod["stock_initialized"]:
                sync_date = prod["stock_sync_date"] or ""
                new_sales = sum(qty for d, qty in entries if d > sync_date)
                if new_sales > 0:
                    new_stock     = max(0, prod["current_stock"] - new_sales)
                    new_sync_date = max(max_upload_date, sync_date)
                    cur.execute(
                        """UPDATE products
                           SET current_stock = ?, stock_sync_date = ?, stock_updated_at = ?
                           WHERE id = ?""",
                        (new_stock, new_sync_date, datetime.now().isoformat(), prod["id"]),
                    )
                    auto_deducted.append(f"{pname} (−{new_sales})")
            else:
                # Track sync date so future deductions are correct once stock is set
                if not prod["stock_sync_date"] or max_upload_date > prod["stock_sync_date"]:
                    cur.execute(
                        "UPDATE products SET stock_sync_date = ? WHERE id = ?",
                        (max_upload_date, prod["id"]),
                    )

        conn.commit()
        conn.close()

        msg = f"Laddade upp {inserted} försäljningsposter"
        if auto_deducted:
            msg += f". Lager automatiskt minskat: {', '.join(auto_deducted)}"
        return jsonify({
            "success":  True,
            "message":  msg,
            "products": list(upload_map.keys()),
        })

    except Exception as exc:
        return jsonify({"error": f"Fel vid uppladdning: {exc}"}), 500


# ─────────────────────────────────────────
# Products & sales
# ─────────────────────────────────────────

@app.get("/api/products")
@require_auth
def get_products():
    conn = get_connection()
    rows = conn.execute(
        "SELECT name, current_stock, stock_initialized, stock_updated_at FROM products ORDER BY name"
    ).fetchall()
    conn.close()
    return jsonify([
        {
            "name":              r["name"],
            "stock":             r["current_stock"],
            "stock_initialized": bool(r["stock_initialized"]),
            "stock_updated_at":  r["stock_updated_at"],
        }
        for r in rows
    ])


@app.get("/api/sales/<product_name>")
@require_auth
def get_sales(product_name):
    conn = get_connection()
    rows = conn.execute(
        """SELECT s.date, SUM(s.quantity) AS quantity
           FROM   sales s
           JOIN   products p ON s.product_id = p.id
           WHERE  p.name = ?
           GROUP  BY s.date
           ORDER  BY s.date
           LIMIT  90""",
        (product_name,),
    ).fetchall()
    conn.close()
    return jsonify([{"date": r["date"], "quantity": r["quantity"]} for r in rows])


# ─────────────────────────────────────────
# Forecast  (days via ?days= param)
# ─────────────────────────────────────────

@app.get("/api/forecast/<product_name>")
@require_auth
def forecast(product_name):
    days   = _days_param(7)
    result = get_forecast(product_name, days=days)
    if result is None:
        return jsonify({"error": "Hittades inte eller för lite data (minst 2 datapunkter)"}), 404
    return jsonify(result)


# ─────────────────────────────────────────
# Recommendations  (days via ?days= param)
# ─────────────────────────────────────────

@app.get("/api/recommendations")
@require_auth
def recommendations():
    days  = _days_param(7)
    conn  = get_connection()
    names = [r["name"] for r in conn.execute("SELECT name FROM products ORDER BY name").fetchall()]
    conn.close()

    results = []
    for name in names:
        data = get_forecast(name, days=days)
        if data:
            results.append({
                "product":           name,
                "status":            data["status"],
                "current_stock":     data["current_stock"],
                "stock_initialized": data["stock_initialized"],
                "order_quantity":    data["order_quantity"],
                "days_stock":        data["days_stock"],
                "total_forecast":    data["total_forecast"],
                "forecast_days":     data["forecast_days"],
            })

    priority = {"urgent": 0, "soon": 1, "unknown": 2, "plan": 3}
    results.sort(key=lambda x: (priority.get(x["status"], 4), -x["order_quantity"]))
    return jsonify(results)


# ─────────────────────────────────────────
# Stock management
# ─────────────────────────────────────────

@app.put("/api/stock/<product_name>")
@require_auth
def update_stock(product_name):
    body = request.get_json(silent=True) or {}
    try:
        stock = int(body.get("stock", 0))
    except (ValueError, TypeError):
        return jsonify({"error": "Ogiltigt lagervärde"}), 400

    now_iso  = datetime.now().isoformat()
    today    = date.today().isoformat()
    conn     = get_connection()
    # Set sync_date to max(today, existing_sync_date) so we don't re-deduct old sales
    conn.execute(
        """UPDATE products
           SET current_stock     = ?,
               stock_initialized = 1,
               stock_updated_at  = ?,
               stock_sync_date   = CASE
                   WHEN stock_sync_date IS NULL OR stock_sync_date < ? THEN ?
                   ELSE stock_sync_date
               END
           WHERE name = ?""",
        (stock, now_iso, today, today, product_name),
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.post("/api/upload-stock")
@require_auth
def upload_stock():
    if "file" not in request.files:
        return jsonify({"error": "Ingen fil uppladdad"}), 400
    f    = request.files["file"]
    name = (f.filename or "").lower()

    try:
        if name.endswith(".csv"):
            df = pd.read_csv(f, encoding="utf-8-sig")
        elif name.endswith((".xlsx", ".xls")):
            df = pd.read_excel(f)
        else:
            return jsonify({"error": "Stödjer bara CSV och Excel"}), 400

        df.columns = df.columns.str.lower().str.strip()
        VARA_NAMES  = {"vara", "produkt", "product", "artikel", "item"}
        LAGER_NAMES = {"lager", "stock", "saldo", "lagerantal", "antal", "quantity"}

        col_vara  = next((c for c in df.columns if c in VARA_NAMES), None)
        col_lager = next((c for c in df.columns if c in LAGER_NAMES), None)
        missing   = [n for n, c in [("vara", col_vara), ("lager", col_lager)] if not c]
        if missing:
            return jsonify({"error": f"Saknar kolumn(er): {', '.join(missing)}"}), 400

        conn             = get_connection()
        updated, skipped = 0, 0
        now_iso          = datetime.now().isoformat()
        today            = date.today().isoformat()

        for _, row in df.iterrows():
            try:
                pname = str(row[col_vara]).strip()
                stock = int(float(str(row[col_lager])))
                if not pname or pname == "nan":
                    continue
            except (ValueError, TypeError):
                skipped += 1
                continue

            result = conn.execute(
                """UPDATE products
                   SET current_stock     = ?,
                       stock_initialized = 1,
                       stock_updated_at  = ?,
                       stock_sync_date   = CASE
                           WHEN stock_sync_date IS NULL OR stock_sync_date < ? THEN ?
                           ELSE stock_sync_date
                       END
                   WHERE name = ?""",
                (stock, now_iso, today, today, pname),
            )
            if result.rowcount > 0:
                updated += 1
            else:
                skipped += 1

        conn.commit()
        conn.close()
        suffix = f", {skipped} hoppades över" if skipped else ""
        return jsonify({
            "success": True,
            "message": f"Uppdaterade lager för {updated} produkt(er){suffix}",
            "updated": updated,
            "skipped": skipped,
        })

    except Exception as exc:
        return jsonify({"error": f"Fel vid uppladdning: {exc}"}), 500


@app.get("/api/stock-template")
@require_auth
def stock_template():
    path = os.path.abspath(
        os.path.join(os.path.dirname(__file__), "..", "lager_exempel.xlsx")
    )
    if not os.path.exists(path):
        return jsonify({"error": "Mallfilen hittades inte"}), 404
    return send_file(path, as_attachment=True, download_name="lager_mall.xlsx")


# ─────────────────────────────────────────
# Export: order recommendations as Excel
# ─────────────────────────────────────────

@app.get("/api/export")
@require_auth
def export_recommendations():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    days  = _days_param(7)
    conn  = get_connection()
    names = [r["name"] for r in conn.execute("SELECT name FROM products ORDER BY name").fetchall()]
    conn.close()

    STATUS_SV   = {"urgent": "Brådskande", "soon": "Snart", "plan": "Planera", "unknown": "Lager okänt"}
    STATUS_FILL = {"Brådskande": "FEE2E2", "Snart": "FEF3C7", "Planera": "DCFCE7", "Lager okänt": "F1F5F9"}

    rows = []
    for name in names:
        data = get_forecast(name, days=days)
        if not data:
            continue
        rows.append({
            "Produkt":               name,
            "Nuv. lager":            data["current_stock"] if data["stock_initialized"] else "Ej angivet",
            f"Prognos {days}d":      data["total_forecast"],
            "Beställ antal":         data["order_quantity"],
            "Status":                STATUS_SV.get(data["status"], data["status"]),
            "Lagerdagar kvar":       data["days_stock"] if data["days_stock"] is not None else "–",
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Beställningslista"

    headers = ["Produkt", "Nuv. lager", f"Prognos {days}d", "Beställ antal", "Status", "Lagerdagar kvar"]
    col_widths = [26, 14, 14, 14, 14, 16]

    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell      = ws.cell(row=1, column=ci, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w

    for ri, row in enumerate(rows, 2):
        vals = [row[h] for h in headers]
        fill_color = STATUS_FILL.get(row["Status"])
        for ci, val in enumerate(vals, 1):
            cell           = ws.cell(row=ri, column=ci, value=val)
            cell.alignment = Alignment(horizontal="left" if ci == 1 else "center")
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)

    # Summary row
    ws.append([])
    ws.append(["Exportdatum:", datetime.now().strftime("%Y-%m-%d %H:%M"), f"Horisont: {days} dagar"])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    date_str = datetime.now().strftime("%Y%m%d")
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"bestallningslista_{date_str}_{days}d.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ─────────────────────────────────────────
# Demo data
# ─────────────────────────────────────────

@app.post("/api/demo")
@require_auth
def demo():
    try:
        generate_demo_data()
        return jsonify({"success": True, "message": "Demo-data genererad (7 produkter, 90 dagar)"})
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(debug=False, host="0.0.0.0", port=port)
